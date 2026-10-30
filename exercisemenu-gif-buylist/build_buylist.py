"""Build the GymVisual GIF buy list for the 20 written workout plans.

Usage: python3 build_buylist.py <use.json> <out.xlsx>
use.json = {canonical exercise: {plan folder: rows}} produced from the plans' workout tables.
"""
import json
import sys
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

use = json.load(open(sys.argv[1]))
out = sys.argv[2]
GV = "https://gymvisual.com/animated-gifs/"


def search(q):
    return "https://gymvisual.com/search?controller=search&orderby=position&orderway=desc&search_query=" + quote_plus(q)


# canonical exercise -> (GymVisual title, link, match, note)
# match: EXACT = same exercise and male model; CLOSE = same movement, small difference noted; SEARCH = no direct page found, use search link
BUY = {
    "Seated Leg Curl (machine)": ("Lever Seated Leg Curl", GV + "2096-lever-seated-leg-curl.html", "EXACT", ""),
    "Cable Triceps Pushdown (rope)": ("Cable Pushdown (with rope attachment)", GV + "1693-cable-pushdown-with-rope-attachment.html", "EXACT", "Also covers 'Triceps Rope Pushdown'."),
    "Leg Press (machine)": ("Sled 45 Leg Press", GV + "2222-sled-45-leg-press.html", "EXACT", ""),
    "Cable Lat Pulldown (wide grip)": ("Cable Bar Lateral Pulldown", GV + "1643-cable-bar-lateral-pulldown.html", "EXACT", "Standard wide-grip bar pulldown; also covers 'Wide-Grip Lat Pulldown'."),
    "Cable Face Pull (rope)": ("Cable Standing Face Pull (with rope)", GV + "14130-cable-standing-face-pull-with-rope.html", "EXACT", ""),
    "Cable Overhead Triceps Extension (rope)": ("Cable Overhead Triceps Extension (rope attachment)", GV + "1687-cable-overhead-triceps-extension-rope-attachment.html", "EXACT", ""),
    "Pull-Up": ("Pull up", GV + "2145-pull-up.html", "EXACT", "Optional extra for 'Weighted Pull-Up': Weighted Pull Up " + GV + "2322-weighted-pull-up.html"),
    "Leg Extension (machine)": ("Lever Leg Extension", GV + "2082-lever-leg-extension.html", "EXACT", ""),
    "Chest-Supported Dumbbell Row (incline bench)": ("Dumbbell Incline Row", GV + "1828-dumbbell-incline-row.html", "EXACT", "Chest on an incline bench."),
    "Seated Cable Row (V-handle)": ("Cable Seated Row with V bar", GV + "6668-cable-seated-row-with-v-bar.html", "EXACT", "Also covers 'Close-Grip Seated Cable Row'."),
    "Cable Kneeling Crunch": ("Cable Kneeling Crunch", GV + "1668-cable-kneeling-crunch.html", "EXACT", ""),
    "Cable One-Arm Lateral Raise": ("Cable Lateral Raise", GV + "1586-cable-lateral-raise.html", "CLOSE", "Check the preview shows one arm; male one-arm version is also listed as 'Cable Single Arm Lateral Raise (VERSION 2) (male)'."),
    "Incline Dumbbell Curl": ("Dumbbell Incline Curl", GV + "1819-dumbbell-incline-curl.html", "EXACT", ""),
    "Reverse Pec Deck (rear delt fly machine)": ("Lever Seated Reverse Fly", GV + "2099-lever-seated-reverse-fly.html", "EXACT", ""),
    "Hanging Knee Raise": ("Hanging Knee Raise (male)", search("Hanging Knee Raise"), "SEARCH", "Listed on GymVisual as 'Hanging Knee Raise (male)' but no direct page was indexed. Do NOT pick 'Assisted Hanging Knee Raise'."),
    "Cable Fly (high to low)": ("Cable Standing Fly", GV + "1720-cable-standing-fly.html", "EXACT", "Standing cable crossover."),
    "45° Back Extension (hyperextension)": ("45 degrees Back Extension", GV + "11527-45-degrees-back-extension.html", "EXACT", "Used in Glute (glute bias) and Core plans."),
    "Chest Dip": ("Chest Dip", GV + "1744-chest-dip.html", "EXACT", "Plan says lean forward for chest — matches."),
    "Dumbbell Romanian Deadlift": ("Dumbbell Romanian Deadlift", GV + "3206-dumbbell-romanian-deadlift.html", "EXACT", ""),
    "Trap Bar Deadlift": ("Trap Bar Deadlift", GV + "2293-trap-bar-deadlift.html", "EXACT", ""),
    "Cable Fly (low to high)": ("Cable Low Fly", GV + "1672-cable-low-fly.html", "EXACT", ""),
    "Dumbbell Walking Lunge": ("Dumbbell Walking Lunges", GV + "3422-dumbbell-walking-lunges.html", "EXACT", "Your owned 'Dumbbell Lunge' GIF is a forward lunge, not a walking lunge."),
    "Plank": ("Front Plank", search("Front Plank"), "SEARCH", "Plain front plank page not indexed. Pick the plain male forearm plank — not 'with Twist', 'Walkout' or 'Weighted'."),
    "Pec Deck Fly (machine)": ("Lever Pec Deck Fly", GV + "2499-lever-pec-deck-fly.html", "EXACT", ""),
    "Cable Straight-Arm Pulldown": ("Cable Straight Arm Pulldown", GV + "1731-cable-straight-arm-pulldown.html", "EXACT", ""),
    "Cable Close-Grip (V-bar) Lat Pulldown": ("Cable Lateral Pulldown with V-bar", GV + "5623-cable-lateral-pulldown-with-v-bar.html", "EXACT", "Covers 'Lat Pulldown (close grip)' and '(neutral grip)'."),
    "Single-Arm Seated Cable Row": ("Cable Seated Single Arm Lats Focused Row", search("Cable Seated Single Arm"), "SEARCH", "Listed on GymVisual but no direct page indexed. Closest direct page: Cable one arm twisting seated row " + GV + "2655-cable-one-arm-twisting-seated-row.html (adds a twist)."),
    "Hack Squat (machine)": ("Sled Hack Squat", GV + "2226-sled-hack-squat.html", "EXACT", ""),
    "Dumbbell Goblet Squat": ("Dumbbell Goblet Squat", GV + "3795-dumbbell-goblet-squat.html", "EXACT", "Plan uses a dumbbell (not kettlebell)."),
    "Dead Bug": ("Dead Bug", GV + "1769-dead-bug.html", "EXACT", ""),
    "Cable Pallof Press": ("Cable horizontal Pallof Press", GV + "2652-cable-horizontal-pallof-press.html", "EXACT", ""),
    "Incline Dumbbell Hammer Curl": ("Dumbbell Incline Hammer Curl", GV + "1821-dumbbell-incline-hammer-curl.html", "EXACT", ""),
    "Hanging Leg Raise": ("Hanging Leg Raise", GV + "1971-hanging-leg-raise.html", "EXACT", ""),
    "Dumbbell Shrug": ("Dumbbell Shrug", GV + "1907-dumbbell-shrug.html", "EXACT", ""),
    "Machine Chest Press": ("Lever Chest Press", GV + "2074-lever-chest-press.html", "EXACT", ""),
    "Cable Biceps Curl (straight bar)": ("Cable Biceps Curl (SZ-bar)", GV + "3471-cable-biceps-curl-sz-bar.html", "CLOSE", "No straight-bar standing cable curl page found; this uses an EZ (SZ) bar attachment — same exercise."),
    "Cable Glute Kickback": ("Cable Kickback (male)", GV + "11282-cable-kickback-male.html", "EXACT", "Glute version (body part: hips). Do NOT pick 'Cable kickback' 2341 — that is the triceps one."),
    "Farmer's Walk": ("Farmers walk", GV + "4220-farmers-walk.html", "EXACT", "Dumbbells."),
    "Close-Grip Dumbbell Press": ("Dumbbell Close grip Press", GV + "1789-dumbbell-close-grip-press.html", "EXACT", ""),
    "Dumbbell Overhead Triceps Extension": ("Dumbbell Seated Triceps Extension", GV + "4223-dumbbell-seated-triceps-extension.html", "EXACT", "One dumbbell held with both hands, as the plan says."),
    "Diamond Push-Up": ("Diamond Push up", GV + "1776-diamond-push-up.html", "EXACT", ""),
    "Band Biceps Curl": ("Band Standard Biceps Curl (male)", GV + "12049-band-standard-biceps-curl-male.html", "EXACT", ""),
    "Band Overhead Triceps Extension": ("Band overhead triceps extension (VERSION 2) (male)", GV + "12144-band-overhead-triceps-extension-version-2-male.html", "EXACT", ""),
    "Band Hammer Curl": ("Band Standing Hammer Curl", GV + "12328-band-standing-hammer-curl.html", "EXACT", ""),
    "Cable Standing Hip Abduction": ("Cable hip abduction (version 2)", GV + "2357-cable-hip-abduction-version-2.html", "EXACT", ""),
    "Dumbbell Reverse Lunge": ("Dumbbell Rear Lunge", GV + "1882-dumbbell-rear-lunge.html", "EXACT", "'Rear lunge' = reverse lunge."),
    "Hip Abduction Machine": ("Lever Seated Hip Abduction", GV + "2094-lever-seated-hip-abduction.html", "EXACT", ""),
    "Dumbbell Hip Thrust": ("Dumbbell Hip Thrust", GV + "10927-dumbbell-hip-thrust.html", "EXACT", ""),
    "Band Side-Lying Hip Abduction": ("Resistance Band Side Lying Hip Abduction", search("Side Lying Hip Abduction"), "SEARCH", "Only a female band version is listed. Male bodyweight alternative: Side Hip Abduction " + GV + "2198-side-hip-abduction.html"),
    "Dumbbell Step-Up": ("Dumbbell Step up", GV + "1932-dumbbell-step-up.html", "EXACT", ""),
    "Single-Leg Calf Raise (bodyweight, on step)": ("Standing Single Leg Calf Raise (On a staircase)", GV + "12400-standing-single-leg-calf-raise-on-a-staircase.html", "EXACT", ""),
    "Side Plank": ("Side Bridge", GV + "2194-side-bridge.html", "CLOSE", "Same exercise ('side bridge' = side plank) but the data sheet says female model. Search 'Side Plank' for a male version if you prefer."),
    "Ab Wheel Rollout": ("Wheel Rollout", GV + "2338-wheel-rollout.html", "EXACT", ""),
    "Cable Woodchop (high to low)": ("Cable Reverse Woodchop (male)", GV + "19083-cable-reverse-woodchop-male.html", "CLOSE", "No high-to-low woodchop GIF found (only a video). This one chops low-to-high — same rotation job. Update the plan wording if you buy this."),
    "Suitcase Carry": ("Dumbbell Suitcase Carry", GV + "6724-dumbbell-suitcase-carry.html", "EXACT", ""),
    "Bird Dog": ("Bird Dog male", GV + "2686-bird-dog-male.html", "EXACT", ""),
    "Reverse Crunch": ("Reverse Crunch (VERSION 2)", GV + "12981-reverse-crunch-version-2.html", "CLOSE", "Plain 'Reverse Crunch' (2164) has a female model; check VERSION 2 preview is male."),
    "Superman": ("Superman (male)", GV + "8422-superman-male.html", "EXACT", "Plan uses a hold — same position."),
    "Bicycle Crunch": ("Bicycle Crunch (male)", GV + "15580-bicycle-crunch-male.html", "EXACT", ""),
    "Dumbbell Fly": ("Dumbbell Fly", GV + "1801-dumbbell-fly.html", "EXACT", ""),
    "Push-Up": ("Push-up", GV + "9830-push-up.html", "EXACT", "Also covers 'Slow-Lowering Push-Up'."),
    "Decline (Feet-Elevated) Push-Up": ("Decline Push up", GV + "1773-decline-push-up.html", "EXACT", ""),
    "Incline Push-Up": ("Incline Push-up (on bench) (male)", GV + "14129-incline-push-up-on-bench-male.html", "EXACT", "Hands on a bench, as the plan says."),
}

OWNED_NOTE = {
    "Barbell Lying Triceps Extension": "Also used for 'EZ-Bar Skull Crusher' (GIF shows a straight bar).",
    "Barbell Preacher Curl": "Used for 'EZ-Bar Preacher Curl' (GIF shows a straight bar).",
    "Barbell Standing Calf Raise": "Used for every 'Standing Calf Raise'.",
    "Barbell Back Squat": "Also covers paused / full-depth / competition-depth squats.",
    "Barbell Bench Press": "Also covers paused and touch-and-go bench press.",
}
UNUSED = ["Dumbbell Lunge (forward lunge)", "EZ Bar Curl"]

need = [k for k in use if not k.startswith("OWN ")]
owned = [k for k in use if k.startswith("OWN ")]
assert set(need) == set(BUY), (set(need) ^ set(BUY))
need.sort(key=lambda k: (-len(use[k]), k))
owned.sort(key=lambda k: (-len(use[k]), k))

plan_name = lambda p: p.split("-", 1)[1].replace("-", " ").title()

wb = Workbook()
hdr_fill = PatternFill("solid", fgColor="161616")
hdr_font = Font(bold=True, color="FFFFFF")
fills = {"EXACT": PatternFill("solid", fgColor="E8F5EE"), "CLOSE": PatternFill("solid", fgColor="FEF3CD"), "SEARCH": PatternFill("solid", fgColor="FDE8E8")}


def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = hdr_fill, hdr_font
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


# Sheet 1: summary
ws = wb.active
ws.title = "Summary"
n_exact = sum(BUY[k][2] == "EXACT" for k in need)
n_close = sum(BUY[k][2] == "CLOSE" for k in need)
n_search = sum(BUY[k][2] == "SEARCH" for k in need)
rows = [
    ["GIFs you already own", 24],
    ["Owned GIFs used in the 20 plans", len(owned)],
    ["Owned GIFs not used in these 20 plans", f"{len(UNUSED)} ({', '.join(UNUSED)})"],
    ["NEW GIFs to buy for the 20 plans", len(need)],
    ["  Direct GymVisual link, exact match", n_exact],
    ["  Direct link, close match (read note)", n_close],
    ["  No direct page indexed — use search link", n_search],
    ["Live exercise pages on exercisemenu.com", "24 (+ 1 duplicate: barbell-lying-triceps-extension-2 — delete it)"],
    ["Price (from GymVisual pages)", "$3.60 per GIF; $0.90 per GIF after 10 items in cart (check at checkout)"],
    ["Which size to download", "720x720 (medium) — same as the GIFs you uploaded before"],
]
sheet(ws, ["Item", "Value"], rows, [44, 90])

# Sheet 2: buy list
ws = wb.create_sheet("Buy List (63)")
rows = []
for i, k in enumerate(need, 1):
    title, link, match, note = BUY[k]
    plans = use[k]
    pr = "1 - High" if len(plans) >= 5 else ("2 - Medium" if len(plans) >= 2 else "3 - Low")
    rows.append([i, pr, k, title, link, match, len(plans), ", ".join(plan_name(p) for p in sorted(plans)), note])
sheet(ws, ["#", "Buy priority", "Exercise in your plans", "GymVisual GIF name", "GymVisual link", "Match", "Used in plans", "Plans", "Note"],
      rows, [5, 12, 34, 40, 60, 9, 8, 60, 60])
for r in range(2, ws.max_row + 1):
    link = ws.cell(r, 5)
    link.hyperlink = link.value
    link.font = Font(color="0563C1", underline="single")
    ws.cell(r, 6).fill = fills[ws.cell(r, 6).value]
ws.auto_filter.ref = ws.dimensions

# Sheet 3: owned GIFs used
ws = wb.create_sheet("Already Owned (22 used)")
rows = []
for k in owned:
    name = k[4:]
    rows.append([name, len(use[k]), ", ".join(plan_name(p) for p in sorted(use[k])), OWNED_NOTE.get(name, "")])
for u in UNUSED:
    rows.append([u, 0, "Not used in these 20 plans", ""])
sheet(ws, ["Owned GIF / exercise page", "Used in plans", "Plans", "Note"], rows, [40, 10, 90, 50])

wb.save(out)
print("saved", out, "need", len(need), "exact", n_exact, "close", n_close, "search", n_search)
